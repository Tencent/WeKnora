#!/usr/bin/env python3
"""Functional build/install checks for WeKnora's bundled runtime. No network I/O."""

def require(condition, detail):
    if not condition:
        raise RuntimeError(str(detail))
import argparse
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
ROOT = Path(__file__).resolve().parent.parent
NAME = ROOT.name
COMMANDS = {'xlsx': ['soffice'], 'docx': ['soffice', 'pdftotext'], 'powerpoint': ['soffice', 'pdftoppm'], 'pdf': ['pdftoppm']}.get(NAME, [])

def run(*args):
    result = subprocess.run(args, check=True, capture_output=True, text=True, timeout=120)
    return result.stdout

def check(directory):
    for command in COMMANDS:
        if not shutil.which(command):
            raise RuntimeError(f'missing runtime command: {command}')
    if NAME == 'xlsx':
        from openpyxl import Workbook, load_workbook
        book = Workbook()
        sheet = book.active
        sheet['A1'], sheet['A2'], sheet['A3'] = (1, 2, '=SUM(A1:A2)')
        source = directory / 'sample.xlsx'
        book.save(source)
        output = directory / 'recalculated.xlsx'
        report = json.loads(run(sys.executable, str(ROOT / 'scripts/xlsx_recalc.py'), str(source), '--out', str(output)))
        require(report.get('recalculated') is True, report)
        require(load_workbook(output, data_only=True).active['A3'].value == 3, "load_workbook(output, data_only=True).active['A3'].value == 3")
    elif NAME == 'docx':
        from docx import Document
        doc = Document()
        doc.add_heading('WeKnora smoke check', 0)
        doc.add_paragraph('Hello document')
        source = directory / 'sample.docx'
        doc.save(source)
        require(Document(source).paragraphs[1].text == 'Hello document', "Document(source).paragraphs[1].text == 'Hello document'")
        run('soffice', '-env:UserInstallation=' + (directory / 'lo-profile').as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(directory), str(source))
        output = directory / 'sample.pdf'
        require(output.read_bytes().startswith(b'%PDF'), "output.read_bytes().startswith(b'%PDF')")
        require('Hello document' in run('pdftotext', str(output), '-'), "'Hello document' in run('pdftotext', str(output), '-')")
    elif NAME == 'powerpoint':
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = 'WeKnora smoke check'
        source = directory / 'sample.pptx'
        presentation.save(source)
        report = json.loads(run(sys.executable, str(ROOT / 'scripts/pptx_render.py'), str(source), '--outdir', str(directory / 'render'), '--dpi', '40'))
        require(report.get('rendered') is True and report.get('files'), report)
        require(all((Path(file).read_bytes().startswith(b'\x89PNG') for file in report['files'])), "all((Path(file).read_bytes().startswith(b'\\x89PNG') for file in report['files']))")
    elif NAME == 'pdf':
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader, PdfWriter
        import pdfplumber
        import pypdfium2
        source = directory / 'sample.pdf'
        c = canvas.Canvas(str(source))
        c.drawString(80, 700, 'Hello PDF')
        c.save()
        require('Hello PDF' in PdfReader(source).pages[0].extract_text(), "'Hello PDF' in PdfReader(source).pages[0].extract_text()")
        with pdfplumber.open(source) as document:
            require('Hello PDF' in document.pages[0].extract_text(), "'Hello PDF' in document.pages[0].extract_text()")
        document = pypdfium2.PdfDocument(str(source))
        page = document[0]
        bitmap = page.render(scale=0.25)
        require(bitmap.width > 0, 'bitmap.width > 0')
        bitmap.close()
        page.close()
        document.close()
        writer = PdfWriter()
        writer.append(source)
        output = directory / 'copy.pdf'
        writer.write(output)
        require(len(PdfReader(output).pages) == 1, 'len(PdfReader(output).pages) == 1')
    elif NAME == 'exploratory-data-analysis':
        import pandas as pd
        frame = pd.DataFrame({'value': [1, 2, None, 4]})
        source = directory / 'sample.csv'
        frame.to_csv(source, index=False)
        require(pd.read_csv(source)['value'].isna().sum() == 1, "pd.read_csv(source)['value'].isna().sum() == 1")
        run(sys.executable, str(ROOT / 'scripts/eda_analyzer.py'), '--help')
    elif NAME == 'statistical-analysis':
        import numpy as np
        from scipy.stats import ttest_ind
        import statsmodels.api as sm
        import pingouin
        result = ttest_ind([1, 2, 3, 4, 5], [4, 5, 6, 7, 8])
        require(0 < result.pvalue < 0.05, '0 < result.pvalue < 0.05')
        model = sm.OLS(np.arange(10) * 2 + 1, sm.add_constant(np.arange(10))).fit()
        require(abs(model.params[1] - 2) < 1e-06, 'abs(model.params[1] - 2) < 1e-06')
    elif NAME == 'scientific-visualization':
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        import seaborn
        import plotly.graph_objects as go
        figure, ax = plt.subplots()
        ax.plot([0, 1, 2], [0, 2, 1])
        for extension in ['png', 'svg', 'pdf']:
            output = directory / ('figure.' + extension)
            figure.savefig(output)
            require(output.stat().st_size > 100, 'output.stat().st_size > 100')
        plt.close(figure)
        require('scatter' in go.Figure(go.Scatter(x=[1], y=[2])).to_json(), "'scatter' in go.Figure(go.Scatter(x=[1], y=[2])).to_json()")
    elif NAME == 'browser':
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.set_content('<label>Name<input aria-label="Name"></label>')
                page.get_by_label('Name').fill('WeKnora')
                require(page.get_by_label('Name').input_value() == 'WeKnora', "page.get_by_label('Name').input_value() == 'WeKnora'")
                require(page.screenshot().startswith(b'\x89PNG'), "page.screenshot().startswith(b'\\x89PNG')")
            finally:
                browser.close()
    else:
        raise RuntimeError('unknown builtin skill')

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--report', action='store_true')
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='weknora-skill-smoke-') as tmp:
        check(Path(tmp))
    if args.report:
        report = ROOT / '.weknora/install-report.json'
        report.parent.mkdir(exist_ok=True)
        report.write_text(json.dumps({'commands': COMMANDS, 'blockers': []}) + '\n')
    print(json.dumps({'ok': True, 'skill': NAME, 'functional_check': 'passed'}))
if __name__ == '__main__':
    main()
