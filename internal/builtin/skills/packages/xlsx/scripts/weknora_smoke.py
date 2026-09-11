#!/usr/bin/env python3
"""Functional checks for the xlsx skill. No network I/O."""

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
NAME = ROOT.name
COMMANDS = ['soffice']


def require(condition, detail):
    if not condition:
        raise RuntimeError(str(detail))


def run(*args):
    return subprocess.run(args, check=True, capture_output=True, text=True, timeout=120).stdout


def check(directory):
    for command in COMMANDS:
        if not shutil.which(command):
            raise RuntimeError(f'missing runtime command: {command}')
    from openpyxl import Workbook, load_workbook
    book = Workbook()
    sheet = book.active
    sheet['A1'], sheet['A2'], sheet['A3'] = (1, 2, '=SUM(A1:A2)')
    source = directory / 'sample.xlsx'
    book.save(source)
    output = directory / 'recalculated.xlsx'
    report = json.loads(run(sys.executable, str(ROOT / 'scripts/xlsx_recalc.py'), str(source), '--out', str(output)))
    require(report.get('recalculated') is True, report)
    require(load_workbook(output, data_only=True).active['A3'].value == 3, "load_workbook(output, data_only=True).active['A3'].value == 3")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--report', action='store_true')
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='weknora-skill-smoke-') as tmp:
        check(Path(tmp))
    if args.report:
        report = ROOT / '.weknora/install-report.json'
        report.parent.mkdir(exist_ok=True)
        report.write_text(json.dumps({'commands': COMMANDS, 'blockers': []}) + '\n')
    print(json.dumps({'ok': True, 'skill': NAME, 'functional_check': 'passed'}))
if __name__ == '__main__':
    main()
