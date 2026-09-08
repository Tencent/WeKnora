from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
from io import BytesIO, StringIO, TextIOWrapper
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).absolute().parents[1]
SCRIPT = ROOT / "scripts" / "build_presentation.py"
sys.path.insert(0, str(SCRIPT.parent))
import build_presentation as builder


def outline():
    return {"title": "Local Runtime", "subtitle": "Synthetic test",
            "slides": [{"type": "bullets", "title": "Execution", "bullets": ["Validate", "Build"]}]}


class BuilderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.workspace = self.root / "workspace"
        self.input = self.workspace / "input"
        self.output = self.workspace / "output"
        self.assets = self.root / "skill" / "assets"
        for path in (self.input, self.output, self.assets):
            path.mkdir(parents=True)
        self.builder = builder.PresentationBuilder(workspace_root=self.workspace, assets_root=self.assets)

    def load(self, value):
        return self.builder.load("-", BytesIO(json.dumps(value).encode("utf-8")))

    def image(self, path, size=(640, 320), image_format="PNG"):
        image = Image.new("RGB", size, "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((20, 40, 200, 80), fill="#007f82")
        draw.rectangle((20, 100, 320, 140), fill="#142d4e")
        image.save(path, format=image_format)
        return path

    def image_deck(self, path):
        return {"title": "Local Chart", "slides": [
            {"type": "image", "title": "Synthetic Fixture", "path": str(path), "caption": "Test data"}
        ]}

    def assert_no_output(self):
        self.assertEqual(list(self.output.iterdir()), [])

    def test_bundle_metadata_and_dependency_pin(self):
        skill = (ROOT / "SKILL.md").read_text(encoding="utf-8")
        self.assertTrue(skill.startswith("---\nname: presentation-builder\n"))
        description = next(line for line in skill.splitlines() if line.startswith("description: "))
        description = json.loads(description.removeprefix("description: "))
        self.assertLess(len(description), 200)
        self.assertIn("PPTX slide deck from structured JSON", description)
        self.assertIn("Use when", description)
        self.assertIn('"skill_name": "presentation-builder"', skill)
        self.assertIn("skill://presentation-builder/SKILL.md", skill)
        self.assertIn("python-pptx==1.0.2", (ROOT / "requirements.txt").read_text())

    def test_sample_zip_structure_slide_count_text_and_size(self):
        demo = builder.PresentationBuilder(workspace_root=self.workspace, assets_root=ROOT / "assets")
        deck = demo.load(ROOT / "assets" / "sample.json")
        result = demo.build(deck, "runtime-demo.pptx")
        path = self.output / "runtime-demo.pptx"
        self.assertEqual(result["path"], str(path))
        self.assertEqual(result["slide_count"], 5)
        self.assertEqual(result["size_bytes"], path.stat().st_size)
        self.assertGreater(result["size_bytes"], 25_000)
        self.assertLess(result["size_bytes"], 150_000)
        with ZipFile(path) as archive:
            self.assertIsNone(archive.testzip())
            self.assertTrue({"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml",
                             "ppt/_rels/presentation.xml.rels", "docProps/core.xml"} <= set(archive.namelist()))
            slides = [name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
            self.assertEqual(len(slides), 5)
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                  "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
            for name in archive.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    ET.fromstring(archive.read(name))
            for name in slides:
                text = ET.fromstring(archive.read(name)).findall(".//a:t", ns)
                self.assertTrue(any(node.text and node.text.strip() for node in text), name)
            presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
            self.assertEqual(len(presentation.findall("p:sldIdLst/p:sldId", ns)), 5)
        prs = Presentation(path)
        self.assertEqual(len(prs.slides), 5)
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=5)
        self.assertEqual(prs.core_properties.title, deck["title"])
        tables = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table]
        self.assertEqual(len(tables), 2)
        self.assertEqual(tables[0].cell(1, 0).text, "/workspace/input")
        for slide in prs.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, prs.slide_width)
                self.assertLessEqual(shape.top + shape.height, prs.slide_height)

    def test_workspace_json_stdin_and_bundled_json(self):
        deck = outline()
        self.assertEqual(self.load(deck), deck)
        for path in (self.workspace / "deck.json", self.input / "deck.json", self.assets / "sample.json"):
            path.write_text(json.dumps(deck), encoding="utf-8")
            self.assertEqual(self.builder.load(path), deck)

    def test_explicit_run_styles_for_browser_office_viewers(self):
        deck = outline()
        deck["slides"].append({"type": "table", "title": "Backend", "headers": ["Name"], "rows": [["Docker"]]})
        result = self.builder.build(deck, "styles.pptx")
        prs = Presentation(result["path"])
        title = next(shape for shape in prs.slides[0].shapes if shape.has_text_frame and shape.text == deck["title"])
        font = title.text_frame.paragraphs[0].runs[0].font
        self.assertEqual(font.size.pt, 36)
        self.assertTrue(font.bold)
        self.assertEqual(str(font.color.rgb), "FFFFFF")
        table = next(shape.table for shape in prs.slides[-1].shapes if shape.has_table)
        header = table.cell(0, 0).text_frame.paragraphs[0].runs[0].font
        self.assertEqual(header.size.pt, 14)
        self.assertTrue(header.bold)
        self.assertEqual(str(header.color.rgb), "FFFFFF")

    def test_non_ascii_text_is_preserved(self):
        deck = outline()
        deck["title"] = "\u672c\u5730\u6f14\u793a Caf\u00e9"
        result = self.builder.build(self.load(deck), "unicode.pptx")
        prs = Presentation(result["path"])
        self.assertTrue(any(shape.has_text_frame and shape.text == deck["title"]
                            for shape in prs.slides[0].shapes))

    def test_valid_maximum_slide_and_text_bounds(self):
        deck = {"title": "t" * 100, "subtitle": "s" * 180, "slides": [
            {"type": "bullets", "title": "h" * 80, "bullets": ["b" * 120] * 6}
        ] * 30}
        result = self.builder.build(self.load(deck), "maximum.pptx")
        self.assertEqual(result["slide_count"], 31)
        self.assertEqual(len(Presentation(result["path"]).slides), 31)
        table = {"title": "Table", "slides": [{"type": "table", "title": "Bounds",
                 "headers": ["h" * 40] * 5, "rows": [["c" * 60] * 5] * 8}]}
        self.load(table)

    def test_invalid_json_and_byte_limit(self):
        bad = [b"", b"{", b"{} trailing", b"\xff", b"\xef\xbb\xbf{}",
               b'{"title":"one","title":"two","slides":[]}', b'{"title":NaN}',
               b'{"title":Infinity}', b'{"title":-Infinity}', b"[" * 2000 + b"]" * 2000,
               b" " * (builder.MAX_INPUT_BYTES + 1)]
        for data in bad:
            with self.subTest(data=data[:40]), self.assertRaises(builder.BuildError):
                self.builder.load("-", BytesIO(data))
        path = self.workspace / "oversized.json"
        path.write_bytes(b" " * (builder.MAX_INPUT_BYTES + 1))
        with self.assertRaises(builder.BuildError):
            self.builder.load(path)
        self.assert_no_output()

    def test_invalid_top_level_schema(self):
        for value in (None, [], True, 42, "text", {}, {"title": "Deck", "slides": []}):
            with self.subTest(value=value), self.assertRaises(builder.BuildError):
                self.load(value)
        for field, value in [("title", None), ("title", " "), ("title", 7), ("title", "x" * 101),
                             ("subtitle", True), ("subtitle", "x" * 181), ("slides", {}),
                             ("slides", [outline()["slides"][0]] * 31), ("unknown", "ignored?")]:
            deck = outline()
            deck[field] = value
            with self.subTest(field=field, value=value), self.assertRaises(builder.BuildError):
                self.load(deck)

    def test_control_characters_and_invalid_unicode(self):
        for character in ("\n", "\r", "\t", "\x00", "\x0b", "\x7f", "\x85", "\u2028", "\u2029", "\ud800", "\ufffe"):
            deck = outline()
            deck["slides"][0]["bullets"] = [f"before{character}after"]
            with self.subTest(character=repr(character)), self.assertRaises(builder.BuildError):
                self.load(deck)

    def test_invalid_slide_schemas(self):
        valid = outline()["slides"][0]
        bad = [None, [], {"type": "chart", "title": "No"}, {"type": [], "title": "No"}]
        for field, value in [("title", ""), ("title", "x" * 81), ("bullets", []),
                             ("bullets", ["x"] * 7), ("bullets", [True]), ("bullets", [""]),
                             ("bullets", ["x" * 121]), ("bullets", "not a list"), ("extra", 1)]:
            slide = deepcopy(valid)
            slide[field] = value
            bad.append(slide)
        bad.append({"title": "Missing type", "bullets": ["one"]})
        for slide in bad:
            with self.subTest(slide=slide), self.assertRaises(builder.BuildError):
                self.load({"title": "Deck", "slides": [slide]})

    def test_invalid_table_schemas(self):
        valid = {"type": "table", "title": "Table", "headers": ["Key", "Value"], "rows": [["A", "B"]]}
        for field, value in [("headers", []), ("headers", ["x"] * 6), ("headers", [""]),
                             ("headers", ["x" * 41]), ("rows", []), ("rows", [["a", "b"]] * 9),
                             ("rows", [["only one"]]), ("rows", [["a", "b", "c"]]),
                             ("rows", ["ab"]), ("rows", [["a", 1]]), ("rows", [["a", "b" * 61]])]:
            slide = deepcopy(valid)
            slide[field] = value
            with self.subTest(field=field, value=value), self.assertRaises(builder.BuildError):
                self.load({"title": "Deck", "slides": [slide]})

    def test_invalid_image_schema(self):
        for field, value in [("path", None), ("path", ""), ("path", "x" * 513), ("caption", "x" * 181)]:
            deck = self.image_deck("assets/chart.png")
            deck["slides"][0][field] = value
            with self.subTest(field=field), self.assertRaises(builder.BuildError):
                self.load(deck)

    def test_input_paths_reject_traversal_outside_roots_and_symlinks(self):
        external = self.root / "external.json"
        external.write_text(json.dumps(outline()))
        prefix = self.root / "workspace-extra"
        prefix.mkdir()
        (prefix / "deck.json").write_text(json.dumps(outline()))
        (self.workspace / "link.json").symlink_to(external)
        (self.workspace / "missing.json").symlink_to(self.root / "missing")
        (self.workspace / "linked-dir").symlink_to(prefix, target_is_directory=True)
        for path in [external, prefix / "deck.json", "deck.json", "https://example.test/deck.json",
                     str(self.workspace) + "/input/../deck.json", self.workspace / "link.json",
                     self.workspace / "missing.json", self.workspace / "linked-dir" / "deck.json"]:
            with self.subTest(path=path), self.assertRaises(builder.BuildError):
                self.builder.load(path)

    def test_input_rejects_fifo_and_directory(self):
        fifo = self.input / "pipe.json"
        os.mkfifo(fifo)
        for path in (fifo, self.input):
            with self.subTest(path=path), self.assertRaisesRegex(builder.BuildError, "regular file"):
                self.builder.load(path)

    def test_input_symlink_swap_before_open_is_rejected(self):
        source = self.workspace / "race.json"
        source.write_text(json.dumps(outline()))
        external = self.root / "external.json"
        external.write_text(json.dumps(outline()))
        original_open = os.open

        def swap_before_open(path, flags, *args, **kwargs):
            if path == source.name:
                source.unlink()
                source.symlink_to(external)
            return original_open(path, flags, *args, **kwargs)

        with patch.object(builder.os, "open", side_effect=swap_before_open), self.assertRaises(builder.BuildError):
            self.builder.load(source)
        self.assert_no_output()

    def test_constructor_roots_cannot_hide_symlinked_ancestors(self):
        alias = self.root / "alias"
        alias.symlink_to(self.workspace, target_is_directory=True)
        unsafe = builder.PresentationBuilder(workspace_root=alias, assets_root=self.assets)
        with self.assertRaises(builder.BuildError):
            unsafe.build(outline(), "no.pptx")
        assets_alias = self.root / "assets-alias"
        assets_alias.symlink_to(self.assets, target_is_directory=True)
        self.image(self.assets / "chart.png")
        unsafe = builder.PresentationBuilder(workspace_root=self.workspace, assets_root=assets_alias)
        with self.assertRaises(builder.BuildError):
            unsafe.build(self.image_deck("assets/chart.png"), "no.pptx")
        self.assert_no_output()

    def test_output_only_accepts_safe_pptx_basenames(self):
        for name in ("", "../escape.pptx", "/tmp/escape.pptx", "output/file.pptx", "a\\b.pptx",
                     "a.pptx\n", "a.pptx\x00", "file.PPTX", ".hidden.pptx", "two words.pptx",
                     "x" * 81 + ".pptx", "file.pptx.zip", "file..pptx", None):
            with self.subTest(name=name), self.assertRaises(builder.BuildError):
                self.builder.build(outline(), name)
        self.assert_no_output()
        self.builder.build(outline(), "A" * 80 + ".pptx")

    def test_existing_files_links_and_directories_are_never_overwritten(self):
        external = self.root / "keep.pptx"
        external.write_bytes(b"original")
        (self.output / "regular.pptx").write_bytes(b"original")
        (self.output / "symlink.pptx").symlink_to(external)
        (self.output / "dangling.pptx").symlink_to(self.root / "not-created")
        os.link(external, self.output / "hardlink.pptx")
        (self.output / "directory.pptx").mkdir()
        for path in list(self.output.iterdir()):
            with self.subTest(name=path.name), self.assertRaises(builder.BuildError):
                self.builder.build(outline(), path.name)
        self.assertEqual(external.read_bytes(), b"original")
        self.assertEqual((self.output / "regular.pptx").read_bytes(), b"original")
        self.assertFalse((self.root / "not-created").exists())
        self.assertEqual(len(list(self.output.iterdir())), 5)

    def test_output_directory_creation_and_symlink_rejection(self):
        self.output.rmdir()
        self.builder.build(outline(), "created.pptx")
        (self.output / "created.pptx").unlink()
        self.output.rmdir()
        outside = self.root / "outside"
        outside.mkdir()
        self.output.symlink_to(outside, target_is_directory=True)
        with self.assertRaises(builder.BuildError):
            self.builder.build(outline(), "no.pptx")
        self.assertEqual(list(outside.iterdir()), [])

    def test_publish_failure_cleans_temporary_file(self):
        for operation in ("link", "fsync"):
            with self.subTest(operation=operation), patch.object(builder.os, operation, side_effect=OSError("disk error")):
                with self.assertRaises(OSError):
                    self.builder.build(outline(), "failed.pptx")
            self.assert_no_output()

    def test_destination_symlink_race_does_not_replace_target(self):
        external = self.root / "keep.pptx"
        external.write_bytes(b"original")
        original_link = os.link

        def link_after_swap(source, destination, **kwargs):
            (self.output / destination).symlink_to(external)
            return original_link(source, destination, **kwargs)

        with patch.object(builder.os, "link", side_effect=link_after_swap), self.assertRaises(builder.BuildError):
            self.builder.build(outline(), "race.pptx")
        self.assertEqual(external.read_bytes(), b"original")
        self.assertTrue((self.output / "race.pptx").is_symlink())
        self.assertEqual(len(list(self.output.iterdir())), 1)

    def test_output_size_limit_does_not_publish(self):
        with patch.object(builder, "MAX_OUTPUT_BYTES", 1), self.assertRaises(builder.BuildError):
            self.builder.build(outline(), "oversize.pptx")
        self.assert_no_output()

    def test_png_and_jpeg_are_embedded_with_aspect_ratio(self):
        png = self.image(self.input / "chart.png")
        jpeg = self.image(self.assets / "chart.jpg", size=(320, 640), image_format="JPEG")
        original = png.read_bytes()
        deck = self.image_deck(png)
        deck["slides"] += self.image_deck("assets/chart.jpg")["slides"]
        deck["slides"] += self.image_deck(jpeg)["slides"]
        result = self.builder.build(self.load(deck), "images.pptx")
        self.assertEqual(result["slide_count"], 4)
        prs = Presentation(result["path"])
        for slide, ratio in zip(list(prs.slides)[1:], (2, 0.5, 0.5)):
            pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
            self.assertEqual(len(pictures), 1)
            self.assertAlmostEqual(pictures[0].width / pictures[0].height, ratio, places=5)
        with ZipFile(result["path"]) as archive:
            self.assertEqual(len([name for name in archive.namelist() if name.startswith("ppt/media/")]), 2)
        self.assertEqual(png.read_bytes(), original)

    def test_image_paths_reject_urls_traversal_and_unapproved_workspace_files(self):
        outside = self.image(self.workspace / "scratch.png")
        linked = self.input / "linked.png"
        linked.symlink_to(outside)
        paths = [outside, linked, "https://example.test/chart.png", "file:///etc/chart.png",
                 "chart.png", "assets/../chart.png", str(self.input) + "/../scratch.png"]
        for path in paths:
            with self.subTest(path=path), self.assertRaises(builder.BuildError):
                self.builder.build(self.image_deck(path), "no.pptx")
        self.assert_no_output()

    def test_image_corruption_format_byte_pixel_and_total_limits(self):
        path = self.input / "bad.png"
        path.write_bytes(b"not an image")
        with self.assertRaises(builder.BuildError):
            self.builder.build(self.image_deck(path), "no.pptx")
        self.image(path, image_format="GIF")
        with self.assertRaises(builder.BuildError):
            self.builder.build(self.image_deck(path), "no.pptx")
        path.write_bytes(b"x" * (builder.MAX_IMAGE_BYTES + 1))
        with self.assertRaises(builder.BuildError):
            self.builder.build(self.image_deck(path), "no.pptx")
        for size in ((4097, 1), (4000, 3001)):
            Image.new("RGB", size).save(path)
            with self.subTest(size=size), self.assertRaises(builder.BuildError):
                self.builder.build(self.image_deck(path), "no.pptx")
        self.image(path)
        deck = self.image_deck(path)
        deck["slides"] *= 2
        with patch.object(builder, "MAX_TOTAL_IMAGE_BYTES", path.stat().st_size), self.assertRaises(builder.BuildError):
            self.builder.build(deck, "no.pptx")
        self.assert_no_output()

    def test_cli_manifest_and_expected_errors(self):
        stdout, stderr = StringIO(), StringIO()
        stdin = TextIOWrapper(BytesIO(json.dumps(outline()).encode()))
        with patch.object(builder, "PresentationBuilder", return_value=self.builder), \
                patch.object(sys, "stdin", stdin), redirect_stdout(stdout), redirect_stderr(stderr):
            status = builder.main(["--input", "-", "--output", "cli.pptx"])
        self.assertEqual(status, 0)
        self.assertEqual(stderr.getvalue(), "")
        self.assertEqual(json.loads(stdout.getvalue())["slide_count"], 2)
        self.assertTrue((self.output / "cli.pptx").is_file())
        for error, expected in ((builder.BuildError("bad input"), 2), (OSError("disk error"), 1),
                                (ImportError("dependency missing"), 1)):
            stdout, stderr = StringIO(), StringIO()
            with patch.object(builder.PresentationBuilder, "load", side_effect=error), \
                    redirect_stdout(stdout), redirect_stderr(stderr):
                status = builder.main(["--input", "-", "--output", "error.pptx"])
            self.assertEqual(status, expected)
            self.assertEqual(stdout.getvalue(), "")
            self.assertIn("error:", stderr.getvalue())
            self.assertNotIn("Traceback", stderr.getvalue())

    def test_real_cli_has_no_root_override_and_does_not_use_environment_roots(self):
        self.assertEqual(builder.PresentationBuilder().workspace_root, Path("/workspace"))
        command = [sys.executable, "-B", str(SCRIPT), "--input", "-", "--output", "no.pptx"]
        for flag in ("--output-root", "--workspace-root", "--assets-root", "--out"):
            result = subprocess.run(command + [flag, str(self.root)], input=b"{}", capture_output=True, timeout=10)
            self.assertEqual(result.returncode, 2)
            self.assertIn(b"unrecognized arguments", result.stderr)
            self.assertEqual(result.stdout, b"")
        with patch.dict(os.environ, {"OUTPUT_ROOT": str(self.root), "WORKSPACE_ROOT": str(self.root),
                                     "WEKNORA_SKILL_DIR": str(self.root)}):
            self.assertEqual(builder.PresentationBuilder().workspace_root, Path("/workspace"))
            self.assertEqual(builder.PresentationBuilder().assets_root, ROOT / "assets")
        result = subprocess.run(command[:-1] + ["../no.pptx"], input=json.dumps(outline()).encode(),
                                capture_output=True, timeout=10)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(result.stdout, b"")
        self.assertIn(b"basename", result.stderr)


if __name__ == "__main__":
    unittest.main()
