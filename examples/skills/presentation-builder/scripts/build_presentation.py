#!/usr/bin/env python3
"""Build a bounded, offline PPTX in the product runtime's output directory."""

import argparse
from contextlib import contextmanager
import errno
from io import BytesIO
import json
import os
from pathlib import Path
import re
import stat
import sys
import uuid


MAX_INPUT_BYTES = 256 * 1024
MAX_IMAGE_BYTES = 4 * 1024 * 1024
MAX_TOTAL_IMAGE_BYTES = 16 * 1024 * 1024
MAX_OUTPUT_BYTES = 32 * 1024 * 1024
ASSETS_ROOT = Path(__file__).absolute().parents[1] / "assets"
OUTPUT_NAME = re.compile(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}\.pptx")


class BuildError(ValueError):
    """Invalid content or an unsafe path; safe to report without a traceback."""


def _object(value, allowed, where):
    if not isinstance(value, dict) or set(value) - set(allowed):
        raise BuildError(f"{where}: expected an object with only {', '.join(allowed)}")


def _text(value, limit, where, empty=False):
    if not isinstance(value, str) or len(value) > limit:
        raise BuildError(f"{where}: expected a string of at most {limit} characters")
    if not empty and not value.strip():
        raise BuildError(f"{where}: text must not be empty")
    if any(ord(c) < 32 or 127 <= ord(c) <= 159 or
           0xD800 <= ord(c) <= 0xDFFF or ord(c) in (0x2028, 0x2029, 0xFFFE, 0xFFFF)
           for c in value):
        raise BuildError(f"{where}: line breaks, control characters, and invalid Unicode are not allowed")


def _array(value, minimum, maximum, where):
    if not isinstance(value, list) or not minimum <= len(value) <= maximum:
        raise BuildError(f"{where}: expected {minimum}-{maximum} items")


def validate_deck(deck):
    """Validate the complete JSON contract using only the standard library."""
    _object(deck, ("title", "subtitle", "slides"), "deck")
    _text(deck.get("title"), 100, "title")
    _text(deck.get("subtitle", ""), 180, "subtitle", empty=True)
    _array(deck.get("slides"), 1, 30, "slides")
    for index, slide in enumerate(deck["slides"]):
        where = f"slides[{index}]"
        if not isinstance(slide, dict):
            raise BuildError(f"{where}: expected a slide object")
        kind = slide.get("type")
        if kind == "bullets":
            _object(slide, ("type", "title", "bullets"), where)
            _array(slide.get("bullets"), 1, 6, f"{where}.bullets")
            for bullet in slide["bullets"]:
                _text(bullet, 120, f"{where}.bullets[]")
        elif kind == "table":
            _object(slide, ("type", "title", "headers", "rows"), where)
            _array(slide.get("headers"), 1, 5, f"{where}.headers")
            for header in slide["headers"]:
                _text(header, 40, f"{where}.headers[]")
            _array(slide.get("rows"), 1, 8, f"{where}.rows")
            width = len(slide["headers"])
            for row in slide["rows"]:
                _array(row, width, width, f"{where}.rows[]")
                for cell in row:
                    _text(cell, 60, f"{where}.rows[][]", empty=True)
        elif kind == "image":
            _object(slide, ("type", "title", "path", "caption"), where)
            _text(slide.get("path"), 512, f"{where}.path")
            _text(slide.get("caption", ""), 180, f"{where}.caption", empty=True)
        else:
            raise BuildError(f"{where}.type: expected bullets, table, or image")
        _text(slide.get("title"), 80, f"{where}.title")


def _unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise BuildError("JSON contains a duplicate object key")
        result[key] = value
    return result


def _reject_constant(value):
    raise BuildError(f"JSON constant {value} is not allowed")


def _absolute_path(value):
    raw = os.fspath(value)
    path = Path(raw)
    if not path.is_absolute() or ".." in path.parts or "\\" in raw or "\x00" in raw:
        raise BuildError("paths must be absolute, without traversal or backslashes")
    return path


@contextmanager
def _directory(path, create_leaf=False):
    # Walk with directory descriptors: no component may redirect through a symlink.
    flags = os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW
    fd = os.open(path.anchor, flags)
    try:
        for index, part in enumerate(path.parts[1:], 1):
            if create_leaf and index == len(path.parts) - 1:
                try:
                    os.mkdir(part, mode=0o700, dir_fd=fd)
                except FileExistsError:
                    pass
            child = os.open(part, flags, dir_fd=fd)
            os.close(fd)
            fd = child
        yield fd
    except OSError as exc:
        if exc.errno in (errno.ELOOP, errno.ENOTDIR):
            raise BuildError("path contains a symlink or a non-directory component") from exc
        raise
    finally:
        os.close(fd)


def _read_local(path, roots, limit):
    path = _absolute_path(path)
    if not any(path.is_relative_to(root) and path != root for root in roots):
        raise BuildError("file is outside the allowed local roots")
    with _directory(path.parent) as parent:
        try:
            fd = os.open(path.name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK,
                         dir_fd=parent)
        except OSError as exc:
            if exc.errno == errno.ELOOP:
                raise BuildError("symlink files are not allowed") from exc
            raise
        try:
            info = os.fstat(fd)
            if not stat.S_ISREG(info.st_mode):
                raise BuildError("input must be a regular file")
            if info.st_size > limit:
                raise BuildError(f"file exceeds the {limit}-byte limit")
            with os.fdopen(fd, "rb", closefd=False) as source:
                data = source.read(limit + 1)
        finally:
            os.close(fd)
    if len(data) > limit:
        raise BuildError(f"file exceeds the {limit}-byte limit")
    return data


class PresentationBuilder:
    def __init__(self, *, workspace_root=Path("/workspace"), assets_root=ASSETS_ROOT):
        # Constructor injection is for tests/embedding; the CLI never accepts roots.
        self.workspace_root = _absolute_path(workspace_root)
        self.assets_root = _absolute_path(assets_root)

    def load(self, source, stdin=None):
        if source == "-":
            stream = sys.stdin.buffer if stdin is None else stdin
            data = stream.read(MAX_INPUT_BYTES + 1)
        else:
            data = _read_local(source, (self.workspace_root, self.assets_root), MAX_INPUT_BYTES)
        if len(data) > MAX_INPUT_BYTES:
            raise BuildError(f"JSON exceeds the {MAX_INPUT_BYTES}-byte limit")
        try:
            deck = json.loads(data.decode("utf-8"), object_pairs_hook=_unique_object,
                              parse_constant=_reject_constant)
        except BuildError:
            raise
        except (ValueError, RecursionError) as exc:
            raise BuildError("input must be a complete UTF-8 JSON object") from exc
        validate_deck(deck)
        return deck

    def _images(self, deck):
        images = {}
        total = 0
        for index, slide in enumerate(deck["slides"]):
            if slide["type"] != "image":
                continue
            from PIL import Image

            raw = slide["path"]
            path = self.assets_root / raw[len("assets/"):] if raw.startswith("assets/") else Path(raw)
            if path.suffix.lower() not in (".png", ".jpg", ".jpeg"):
                raise BuildError("images must be local PNG or JPEG files")
            blob = _read_local(path, (self.workspace_root / "input", self.assets_root), MAX_IMAGE_BYTES)
            total += len(blob)
            if total > MAX_TOTAL_IMAGE_BYTES:
                raise BuildError("images exceed the 16 MiB total limit")
            try:
                with Image.open(BytesIO(blob)) as image:
                    width, height = image.size
                    if image.format not in ("PNG", "JPEG"):
                        raise BuildError("image content must be PNG or JPEG")
                    if not (0 < width <= 4096 and 0 < height <= 4096 and width * height <= 12_000_000):
                        raise BuildError("image exceeds the dimension or pixel limit")
                    image.load()
            except (OSError, ValueError, Image.DecompressionBombError) as exc:
                if isinstance(exc, BuildError):
                    raise
                raise BuildError("image is corrupt, unsupported, or exceeds pixel limits") from exc
            images[index] = (blob, width, height)
        return images

    def build(self, deck, output_name):
        if not isinstance(output_name, str) or not OUTPUT_NAME.fullmatch(output_name):
            raise BuildError("output must be a basename such as presentation.pptx (letters, digits, _ or -)")
        validate_deck(deck)
        payload = _render(deck, self._images(deck))
        if len(payload) > MAX_OUTPUT_BYTES:
            raise BuildError("PPTX exceeds the 32 MiB output limit")
        output_root = self.workspace_root / "output"
        with _directory(output_root, create_leaf=True) as directory:
            temporary = f".presentation-{uuid.uuid4().hex}.tmp"
            fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW,
                         0o600, dir_fd=directory)
            try:
                with os.fdopen(fd, "wb") as target:
                    target.write(payload)
                    target.flush()
                    os.fsync(target.fileno())
                # Publish a complete file without replacing any existing name, even a symlink.
                try:
                    os.link(temporary, output_name, src_dir_fd=directory,
                            dst_dir_fd=directory, follow_symlinks=False)
                except FileExistsError as exc:
                    raise BuildError("output already exists; choose a new basename") from exc
            finally:
                os.unlink(temporary, dir_fd=directory)
        return {"path": str(output_root / output_name),
                "slide_count": len(deck["slides"]) + 1, "size_bytes": len(payload)}


def _render(deck, images):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = deck["title"]
    prs.core_properties.subject = deck.get("subtitle", "")
    prs.core_properties.author = "Presentation Builder"
    prs.core_properties.keywords = ""
    prs.core_properties.comments = ""
    navy, teal, ink, muted = "142D4E", "007F82", "24364B", "53657B"

    def fill(shape, color):
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)

    def rectangle(slide, x, y, width, height, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                       Inches(width), Inches(height))
        fill(shape, color)
        shape.line.fill.background()

    def text(slide, value, x, y, width, height, size, color=ink, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.text = value
        # Lightweight Office viewers may ignore paragraph-default run properties.
        for font in [paragraph.font, *(run.font for run in paragraph.runs)]:
            font.name = "Aptos"
            font.size = Pt(size)
            font.bold = bold
            font.color.rgb = RGBColor.from_string(color)
        paragraph.line_spacing = 1.0
        paragraph.space_after = Pt(0)
        return box

    def slide_base(title, number):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        rectangle(slide, 0, 0, 13.333333, 0.12, teal)
        text(slide, title, 0.6, 0.42, 12.1, 1.08, 28, navy, True)
        text(slide, f"{number:02d} / {len(deck['slides']) + 1:02d}",
             11.8, 7.04, 1, 0.24, 10, muted)
        return slide

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = RGBColor.from_string(navy)
    rectangle(cover, 0.65, 1.35, 1.1, 0.09, "36C5BE")
    text(cover, deck["title"], 0.65, 1.8, 11.95, 2.45, 36, "FFFFFF", True)
    text(cover, deck.get("subtitle", ""), 0.65, 4.6, 11.9, 1.4, 21, "D8E6F1")
    text(cover, f"01 / {len(deck['slides']) + 1:02d}", 11.8, 7.04, 1, 0.24, 10, "D8E6F1")

    for index, spec in enumerate(deck["slides"]):
        slide = slide_base(spec["title"], index + 2)
        if spec["type"] == "bullets":
            for row, bullet in enumerate(spec["bullets"]):
                y = 1.65 + row * 0.86
                rectangle(slide, 0.65, y + 0.13, 0.09, 0.09, teal)
                text(slide, bullet, 0.95, y, 11.7, 0.8, 18)
        elif spec["type"] == "table":
            rows = [spec["headers"]] + spec["rows"]
            table = slide.shapes.add_table(len(rows), len(spec["headers"]),
                                           Inches(0.6), Inches(1.65), Inches(12.1), Inches(5.1)).table
            table.first_row = True
            for r, values in enumerate(rows):
                for c, value in enumerate(values):
                    cell = table.cell(r, c)
                    cell.text = value
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(
                        navy if r == 0 else ("EDF3F7" if r % 2 else "F7FAFC"))
                    cell.margin_left = cell.margin_right = Inches(0.09)
                    cell.margin_top = cell.margin_bottom = Inches(0.04)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.text_frame.word_wrap = True
                    cell.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                    for paragraph in cell.text_frame.paragraphs:
                        for font in [paragraph.font, *(run.font for run in paragraph.runs)]:
                            font.name = "Aptos"
                            font.size = Pt(12 if len(spec["headers"]) > 3 else 14)
                            font.bold = r == 0
                            font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else ink)
                        paragraph.line_spacing = 1.0
                        paragraph.space_after = Pt(0)
        else:
            blob, width, height = images[index]
            scale = min(12.1 / width, 4.45 / height)
            w, h = width * scale, height * scale
            slide.shapes.add_picture(BytesIO(blob), Inches(0.6 + (12.1 - w) / 2),
                                     Inches(1.6 + (4.45 - h) / 2), Inches(w), Inches(h))
            text(slide, spec.get("caption", ""), 0.6, 6.2, 12.1, 0.65, 15, muted)

    output = BytesIO()
    prs.save(output)
    return output.getvalue()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, allow_abbrev=False)
    parser.add_argument("--input", required=True, help="absolute JSON path or - for stdin")
    parser.add_argument("--output", required=True, help="new PPTX basename in /workspace/output")
    args = parser.parse_args(argv)
    try:
        builder = PresentationBuilder()
        result = builder.build(builder.load(args.input), args.output)
    except BuildError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    except ImportError:
        print("error: runtime dependencies unavailable; provision requirements.txt in the skill environment",
              file=sys.stderr)
        return 1
    except OSError as exc:
        print(f"error: filesystem operation failed: {exc}", file=sys.stderr)
        return 1
    print(json.dumps(result))
    return 0


if __name__ == "__main__":
    sys.exit(main())
