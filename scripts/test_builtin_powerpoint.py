#!/usr/bin/env python3
"""Exercise native PPT generation and Chinese/Latin PDF conversion in the image."""
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
from pptx import Presentation

SCRIPTS = Path(os.environ.get('WEKNORA_PPTX_SCRIPTS', '/opt/weknora/builtin/skills/powerpoint/scripts'))

class PowerpointCreationTests(unittest.TestCase):
    def test_editable_chart_and_mixed_text_survive_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / 'mixed.pptx'
            env = dict(os.environ, NODE_PATH=str(SCRIPTS.parent / 'node_modules'))
            subprocess.run(['node', str(SCRIPTS / 'weknora_example.cjs'), str(deck)], check=True, env=env)
            prs = Presentation(str(deck))
            self.assertEqual(len(prs.slides), 2)
            self.assertTrue(any(shape.has_chart for shape in prs.slides[1].shapes))
            self.assertGreater(sum(shape.has_text_frame for slide in prs.slides for shape in slide.shapes), 3)
            subprocess.run(['soffice', f'-env:UserInstallation=file://{tmp}/lo', '--headless', '--convert-to', 'pdf', '--outdir', tmp, str(deck)], check=True, capture_output=True)
            text = subprocess.check_output(['pdftotext', '-layout', str(deck.with_suffix('.pdf')), '-'], text=True)
            for value in ['English', '中文', '示例', 'Office skills']:
                self.assertIn(value, text)
            self.assertNotIn('[object Object]', text)
            fonts = subprocess.check_output(['pdffonts', str(deck.with_suffix('.pdf'))], text=True)
            self.assertIn('NotoSansCJK', fonts)

if __name__ == '__main__': unittest.main()
