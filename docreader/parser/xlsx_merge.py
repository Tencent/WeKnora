"""Fill merged cell values before pandas reads an XLSX workbook."""

from __future__ import annotations

import logging
import os
import re
import zipfile
from io import BytesIO

logger = logging.getLogger(__name__)

# --- Percentage number formats (opt-in) --------------------------------------
# Excel keeps a percentage as the stored fraction plus a display format: the
# cell holds 0.61123366 and "0.00%" renders it as 61.12%. openpyxl and pandas
# read only the fraction, so a value the user sees as 61.12% reaches the
# pipeline as 0.6112 — a 100x error that a model will copy verbatim into a
# report. When this switch is on, percentage cells are materialized as the
# literal text the spreadsheet shows ("61.12%") instead.
#
# It is off by default on purpose: the cell changes from a number into a
# string, so anything that sorts, aggregates or computes with it sees different
# input. With it off, every workbook is read exactly as before. A single upload
# can override it through parser_engine_overrides.xlsx_percent_as_text.
PERCENT_AS_TEXT_ENV = "DOCREADER_XLSX_PERCENT_AS_TEXT"

# Only percentages are materialized. Currency and thousands separators change
# how a number is displayed, not what it means ("1,234,567" and 1234567 are the
# same number), so rewriting them would churn cells for no semantic gain, and
# date formats are already converted to datetimes by openpyxl.

# --- Same-row hint fallback (opt-in, off by default) -------------------------
# Workbooks occasionally leave one cell of a percentage row without a
# percentage format (a real financial statement had 折旧率 0.0592 formatted as
# "0.00_ ;[Red]\-0.00\ "). When a row carries at least this many explicitly
# percentage-formatted cells, an unformatted 0 < |value| < 1 number in that row
# is treated as a percentage too. Only consulted while percent materialization
# is on; 0 disables the fallback.
ROW_HINT_MIN_ENV = "DOCREADER_XLSX_PCT_ROW_HINT_MIN"
_DEFAULT_ROW_HINT_MIN = 0

# ---- 合并区面积闸门（本补丁新增）----
# openpyxl 在 load_workbook 阶段就按合并区面积创建/清理单元格：一个 4.9 KB 的
# xlsx 只要声明 A1:P100000（1.6M 格），加载就要 ~15.6s / ~576MB；声明
# A1:XFD1048576 会直接把 docreader 卡到不可用。闸门必须在交给 openpyxl
# **之前**生效，所以先扫 sheet XML、把面积超限的 <mergeCell> 摘掉。
MERGE_FILL_LIMIT_ENV = "DOCREADER_XLSX_MAX_MERGE_CELLS"
_DEFAULT_MERGE_FILL_LIMIT = 200_000  # 单条合并区上限（格）
MERGE_TOTAL_LIMIT_ENV = "DOCREADER_XLSX_MAX_MERGE_TOTAL_CELLS"
_DEFAULT_MERGE_TOTAL_LIMIT = 2_000_000  # 单文件所有合并区面积之和上限（格）

_MERGE_CELL_TAG = re.compile(r'<mergeCell\b[^>]*\bref="([A-Za-z]+\d+:[A-Za-z]+\d+)"[^>]*/>')
_MERGE_CELLS_COUNT = re.compile(r'(<mergeCells\b[^>]*\bcount=")(\d+)(")')
_SHEET_XML = re.compile(r"^xl/worksheets/[^/]+\.xml$")


def _env_positive_int(name: str, default: int) -> int:
    raw = os.environ.get(name, "").strip()
    if not raw:
        return default
    try:
        value = int(raw)
    except ValueError:
        logger.warning("XLSX pre-parse: 环境变量 %s=%r 不是整数，忽略，用默认值 %d", name, raw, default)
        return default
    if value <= 0:
        logger.warning("XLSX pre-parse: 环境变量 %s=%r 非正数，忽略，用默认值 %d", name, raw, default)
        return default
    return value


def _merge_area(ref: str) -> int:
    """合并区覆盖的单元格数；无法解析返回 0（交给 openpyxl 自己报错）。"""
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    try:
        first, last = ref.split(":")
        col1, row1 = coordinate_from_string(first.strip().upper())
        col2, row2 = coordinate_from_string(last.strip().upper())
        return (column_index_from_string(col2) - column_index_from_string(col1) + 1) * (
            row2 - row1 + 1
        )
    except Exception:
        return 0


def _guard_merge_ranges(content: bytes) -> bytes:
    """摘掉面积超限的合并区声明，避免 openpyxl 在加载阶段被撑爆。

    只改写 xl/worksheets/*.xml；没有超限声明时原样返回（正常文件零开销）。
    """
    per_limit = _env_positive_int(MERGE_FILL_LIMIT_ENV, _DEFAULT_MERGE_FILL_LIMIT)
    total_limit = _env_positive_int(MERGE_TOTAL_LIMIT_ENV, _DEFAULT_MERGE_TOTAL_LIMIT)

    src = zipfile.ZipFile(BytesIO(content))
    rewritten: dict[str, bytes] = {}
    dropped_ranges = 0
    dropped_area = 0
    budget = total_limit
    try:
        for item in src.infolist():
            if not _SHEET_XML.match(item.filename):
                continue
            text = src.read(item.filename).decode("utf-8")
            if "mergeCell" not in text:
                continue
            changed = False
            pieces = []
            cursor = 0
            kept = 0
            for match in _MERGE_CELL_TAG.finditer(text):
                ref = match.group(1)
                area = _merge_area(ref)
                if area > per_limit or area > budget:
                    dropped_ranges += 1
                    dropped_area += area
                    pieces.append(text[cursor:match.start()])
                    cursor = match.end()
                    changed = True
                    continue
                budget -= area
                kept += 1
            if changed:
                pieces.append(text[cursor:])
                new_text = "".join(pieces)
                new_text = _MERGE_CELLS_COUNT.sub(
                    lambda m: f"{m.group(1)}{kept}{m.group(3)}", new_text, count=1
                )
                rewritten[item.filename] = new_text.encode("utf-8")
    finally:
        src.close()

    if not rewritten:
        return content

    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        with zipfile.ZipFile(BytesIO(content)) as original:
            for item in original.infolist():
                data = rewritten.get(item.filename)
                dst.writestr(item, data if data is not None else original.read(item.filename))
    logger.warning(
        "XLSX pre-parse: 摘掉 %d 个超限合并区（合计声明 %d 格；单条上限 %d，文件上限 %d）",
        dropped_ranges,
        dropped_area,
        per_limit,
        total_limit,
    )
    return out.getvalue()


# Only materialize a known numeric percentage pattern. Excel literals, conditions,
# scaling and arbitrary decorations require a full format renderer; leave those
# values untouched rather than guessing that every '%' means multiply by 100.
_PERCENT_PATTERN = re.compile(
    r"(?P<sign>-?)(?P<open>\()?(?P<group>#,##)?0"
    r"(?:\.(?P<decimals>0+))?%(?P<close>\))?"
)
_FORMAT_COLOR = re.compile(
    r"\[(?:Black|Blue|Cyan|Green|Magenta|Red|White|Yellow|Color[0-9]+)\]", re.I
)


def _env_bool(name: str, default: bool) -> bool:
    value = os.environ.get(name)
    if value is None or not str(value).strip():
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on"}


def _env_int(name: str, default: int) -> int:
    try:
        return int(str(os.environ.get(name, default)).strip())
    except (TypeError, ValueError):
        return default


def percent_as_text_default() -> bool:
    """Deployment-wide default for materializing percentage formats.

    Read per call (not at import time) so the value can be flipped without a
    module reload; a per-upload override still wins over it.
    """
    return _env_bool(PERCENT_AS_TEXT_ENV, False)


def _format_percentage(value, number_format: str):
    """Render supported Excel percentage formats, leaving other formats intact."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if not number_format or "%" not in number_format:
        return None
    # Quoted/escaped percentages are literal text, not percentage operators.
    # Reject unsupported conditions across all sections before selecting one.
    cleaned = _FORMAT_COLOR.sub("", number_format)
    if any(char in cleaned for char in ('"', "\\", "_", "*", "[", "]")):
        return None
    sections = cleaned.split(";")
    if len(sections) > 4:
        return None
    section_index = 1 if value < 0 and len(sections) >= 2 else 0
    if value == 0 and len(sections) >= 3:
        section_index = 2
    match = _PERCENT_PATTERN.fullmatch(sections[section_index].strip())
    if not match or bool(match["open"]) != bool(match["close"]):
        return None
    decimals = len(match["decimals"] or "")
    # Excel's negative section formats the magnitude; a single section instead
    # supplies the minus sign automatically.
    amount = abs(value) if section_index == 1 else value
    grouping = "," if match["group"] else ""
    try:
        number = f"{float(amount) * 100:{grouping}.{decimals}f}"
        return f"{match['sign']}{match['open'] or ''}{number}%{match['close'] or ''}"
    except (TypeError, ValueError, OverflowError):
        return None


def _materialize_percentage_cells(ws, row_hint_min: int) -> tuple[int, int]:
    """Rewrite the percentage-formatted numbers of *ws* as their display text.

    Must run before merged ranges are filled: only the master cell of a merge
    carries a number_format while the covered cells are General, so filling
    first would copy a bare fraction into rows 2..n and leave it unconverted —
    the same row would then show both "0.55%" and 0.005509433890047721.

    Returns ``(materialized, hinted)`` cell counts.
    """
    row_hints: dict = {}
    if row_hint_min > 0:
        # Count the explicit percentage formats before converting anything:
        # the conversion below resets number_format to General.
        for row in ws.iter_rows():
            count = sum(
                1
                for cell in row
                if type(cell.value) in (int, float)  # noqa: E721 - excludes bool
                and _format_percentage(cell.value, cell.number_format) is not None
            )
            if count:
                row_hints[row[0].row] = count

    materialized = 0
    hinted = 0
    for row in ws.iter_rows():
        hint = row_hints.get(row[0].row, 0)
        for cell in row:
            if type(cell.value) not in (int, float):  # noqa: E721 - excludes bool
                continue
            text = _format_percentage(cell.value, cell.number_format)
            if (
                text is None
                and "%" not in str(cell.number_format or "")
                and hint >= row_hint_min > 0
                and 0 < abs(cell.value) < 1
            ):
                # Fallback for a cell that lost its percentage format while its
                # row kept one. Logged because it rewrites an unformatted value.
                text = f"{float(cell.value) * 100:.2f}%"
                hinted += 1
                logger.info(
                    "XLSX pre-parse: %s!%s%d has no percentage format but its row "
                    "does; reading it as %.2f%% (stored value %r)",
                    ws.title,
                    cell.column_letter,
                    cell.row,
                    float(cell.value) * 100,
                    cell.value,
                )
            if text is None:
                continue
            cell.value = text
            cell.number_format = "General"
            materialized += 1
    return materialized, hinted


def fill_merged_cells_xlsx(
    content: bytes, percent_as_text: bool | None = None
) -> bytes:
    """Unmerge ranges and copy the master cell value into every covered cell.

    openpyxl only stores values on the top-left cell of a merge; pandas then
    sees NaN in the rest. Filling makes row-wise RAG chunks retain context.

    ``percent_as_text`` also materializes percentage-formatted numbers as the
    text the spreadsheet displays ("0.6112" -> "61.12%"), because the stored
    fraction is 100x off from what a reader sees. ``None`` (the default) uses
    the ``DOCREADER_XLSX_PERCENT_AS_TEXT`` deployment default, which is off, so
    the default path is unchanged.
    """
    if percent_as_text is None:
        percent_as_text = percent_as_text_default()

    if not zipfile.is_zipfile(BytesIO(content)):
        return content

    # 必须在 load_workbook 之前：openpyxl 加载时就会按合并区面积展开。
    content = _guard_merge_ranges(content)

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content), data_only=True)
    changed = False
    materialized_cells = 0
    hinted_cells = 0
    skipped_ranges = 0
    row_hint_min = (
        _env_int(ROW_HINT_MIN_ENV, _DEFAULT_ROW_HINT_MIN) if percent_as_text else 0
    )

    for ws in wb.worksheets:
        if percent_as_text:
            # Before the merge fill below — see _materialize_percentage_cells.
            materialized, hinted = _materialize_percentage_cells(ws, row_hint_min)
            materialized_cells += materialized
            hinted_cells += hinted
            if materialized:
                changed = True

        if not ws.merged_cells.ranges:
            continue
        for merge_range in list(ws.merged_cells.ranges):
            area = (merge_range.max_row - merge_range.min_row + 1) * (
                merge_range.max_col - merge_range.min_col + 1
            )
            if area > _env_positive_int(MERGE_FILL_LIMIT_ENV, _DEFAULT_MERGE_FILL_LIMIT):
                # 兜底：XML 闸门没能摘掉它（非常规写法），这里再挡一次。
                skipped_ranges += 1
                logger.warning(
                    "XLSX pre-parse: %s!%s 合并区 %s 覆盖 %d 格，超过单条上限，跳过填充（保留左上角值）",
                    ws.title,
                    merge_range.coord,
                    area,
                )
                continue
            master_value = ws.cell(merge_range.min_row, merge_range.min_col).value
            ws.unmerge_cells(str(merge_range))
            for row in range(merge_range.min_row, merge_range.max_row + 1):
                for col in range(merge_range.min_col, merge_range.max_col + 1):
                    ws.cell(row, col).value = master_value
            changed = True

    if not changed:
        return content

    # Charts are not part of the cell text extracted by DocReader. Remove
    # dedicated chart sheets through the public workbook API, then clear
    # embedded worksheet charts (openpyxl has no public removal API for them)
    # before saving this temporary workbook.
    active_sheet = wb.active
    for chart_sheet in list(wb.chartsheets):
        wb.remove(chart_sheet)
    if active_sheet in wb.worksheets:
        wb.active = active_sheet
    elif wb.worksheets:
        wb.active = wb.worksheets[0]
    for ws in wb.worksheets:
        ws._charts.clear()

    out = BytesIO()
    wb.save(out)
    if materialized_cells:
        logger.info(
            "Materialized %d percentage cell(s) in XLSX before parse "
            "(%d recovered by the same-row hint fallback, threshold=%d)",
            materialized_cells,
            hinted_cells,
            row_hint_min,
        )
    logger.info("Filled merged cells in XLSX before parse")
    return out.getvalue()
