"""Exercise saved PPTM/XLSM content through both public parser entries."""

import io
import unittest
import zipfile
from types import SimpleNamespace
from unittest.mock import patch

from lxml import etree
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from docreader.parser.parser import Parser
from docreader.parser.markitdown_parser import StdMarkitdownParser
from docreader.parser.registry import registry


def macro_package(data, old_type, new_type, macro_part):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(output, "w") as dst:
        for info in src.infolist():
            part = src.read(info)
            if info.filename == "[Content_Types].xml":
                part = part.replace(old_type.encode(), new_type.encode())
            dst.writestr(info, part)
        dst.writestr(macro_part, b"inert VBA preservation marker")
    return output.getvalue()


def make_pptm():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = (
        "PPTM saved text"
    )
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "Item"
    table.cell(1, 0).text = "Table sentinel"
    table.cell(1, 1).text = "12345"
    picture = io.BytesIO()
    Image.new("RGB", (128, 96), "blue").save(picture, format="PNG")
    picture.seek(0)
    slide.shapes.add_picture(picture, Inches(1), Inches(4))
    output = io.BytesIO()
    deck.save(output)
    return macro_package(
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
        "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
        "ppt/vbaProject.bin",
    )


def make_xlsm():
    book = Workbook()
    sheet = book.active
    sheet.title = "Saved values"
    sheet.append(["Label", "Value"])
    sheet.append(["Cached formula", "=12345+1"])
    sheet.append(["Uncached formula", "=12345+2"])
    sheet.append(["Saved text", "XLSM sentinel"])
    sheet.append(["Uncached VBA function", "=MacroFunction(12345)"])
    output = io.BytesIO()
    book.save(output)
    cached = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as src, zipfile.ZipFile(
        cached, "w"
    ) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = etree.fromstring(data)
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                root.find(".//s:c[@r='B2']/s:v", ns).text = "12346"
                data = etree.tostring(root)
            dst.writestr(info, data)
    return macro_package(
        cached.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        "xl/vbaProject.bin",
    )


class MacroOfficeParserTest(unittest.TestCase):
    def test_pptm_recovers_unresolved_image_references(self):
        with patch.object(
            StdMarkitdownParser,
            "_convert_markitdown",
            return_value=SimpleNamespace(
                text_content="PPTM image\n![picture](image1.png)"
            ),
        ):
            result = StdMarkitdownParser(file_type="pptm").parse_into_text(make_pptm())
        self.assertEqual(len(result.images), 1)
        self.assertNotIn("](image1.png)", result.content)
        for ref in result.images:
            self.assertIn(ref, result.content)

    def test_pptm_saved_text_tables_and_images(self):
        data = make_pptm()
        for engine in ("builtin", "markitdown"):
            with self.subTest(engine=engine), patch(
                "docreader.parser.ppt_convert.convert_ppt_to_pptx_bytes",
                side_effect=AssertionError("PPTM must not require LibreOffice"),
            ):
                result = Parser().parse_file(
                    "sample.pptm", ".PPTM", data, parser_engine=engine
                )
                for expected in ("PPTM saved text", "Table sentinel", "12345"):
                    self.assertIn(expected, result.content)
                self.assertTrue(result.images)
                for ref in result.images:
                    self.assertIn(ref, result.content)

    def test_xlsm_reads_saved_values_without_recalculating_formulas(self):
        data = make_xlsm()
        for engine in ("builtin", "markitdown"):
            with self.subTest(engine=engine), patch(
                "docreader.parser.excel_convert.convert_excel_to_xlsx_bytes",
                side_effect=AssertionError("XLSM must not require LibreOffice"),
            ):
                result = Parser().parse_file(
                    "sample.xlsm", ".XLSM", data, parser_engine=engine
                )
                for expected in (
                    "XLSM sentinel",
                    "Cached formula",
                    "12346",
                    "Uncached formula",
                    "Uncached VBA function",
                ):
                    self.assertIn(expected, result.content)
                self.assertNotIn("12347", result.content)
                self.assertNotIn("MacroFunction", result.content)

    def test_both_engines_advertise_macro_formats(self):
        for engine in ("builtin", "markitdown"):
            listed = next(e for e in registry.list_engines() if e["name"] == engine)
            for ext in ("pptm", "xlsm"):
                self.assertIn(ext, listed["file_types"])


if __name__ == "__main__":
    unittest.main()
