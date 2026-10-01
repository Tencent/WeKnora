import io
import unittest
from unittest.mock import Mock, patch

from docx import Document
from docx.enum.section import WD_SECTION_START

from docreader.parser.docx2_parser import Docx2Parser
from docreader.parser.docx_parser import DocxParser
from docreader.parser.markitdown_parser import MarkitdownParser


class DocxHeadersTest(unittest.TestCase):
    def make_document(self):
        document = Document()
        section = document.sections[0]
        section.header.paragraphs[0].text = "Header identifier 3849"
        section.different_first_page_header_footer = True
        section.first_page_header.paragraphs[0].text = "First page identifier"
        document.settings.odd_and_even_pages_header_footer = True
        section.even_page_header.paragraphs[0].text = "Even page identifier"
        table = section.header.add_table(rows=1, cols=2, width=section.page_width)
        table.cell(0, 0).text = "Header table key"
        table.cell(0, 1).text = "Header table value"
        document.add_paragraph("Body text")
        document.add_section(WD_SECTION_START.NEW_PAGE)
        document.sections[-1].first_page_header.is_linked_to_previous = False
        document.sections[-1].different_first_page_header_footer = False
        document.sections[-1].first_page_header.paragraphs[0].text = "Inactive header"
        document.sections[-1].header.is_linked_to_previous = False
        document.sections[-1].header.paragraphs[0].text = "Second section header"
        document.add_section(WD_SECTION_START.NEW_PAGE)
        document.sections[-1].footer.paragraphs[0].text = "Footer text"
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def test_headers_are_included_when_enabled(self):
        for parser_cls in (Docx2Parser, MarkitdownParser):
            with self.subTest(parser=parser_cls.__name__):
                parsed = parser_cls(
                    file_type="docx", docx_include_headers="true"
                ).parse_into_text(self.make_document())
                self.assertIn("Header identifier 3849", parsed.content)
                self.assertIn("Body text", parsed.content)
                for text in (
                    "First page identifier", "Even page identifier",
                    "Header table key", "Header table value", "Second section header",
                ):
                    self.assertEqual(parsed.content.count(text), 1)
                self.assertNotIn("Inactive header", parsed.content)
                self.assertNotIn("Footer text", parsed.content)

    def test_headers_remain_excluded_by_default_or_when_disabled(self):
        for parser_cls in (Docx2Parser, MarkitdownParser):
            for options in (
                {}, {"docx_include_headers": "false"}, {"docx_include_headers": False},
            ):
                with self.subTest(parser=parser_cls.__name__, options=options):
                    text = parser_cls(file_type="docx", **options).parse_into_text(
                        self.make_document()
                    ).content
                    self.assertIn("Body text", text)
                    self.assertNotIn("Header identifier 3849", text)

    def test_python_docx_fallback_keeps_headers(self):
        with patch.dict(
            DocxParser._parse_body.__globals__,
            {"Docx": Mock(side_effect=RuntimeError("fallback"))},
        ):
            text = DocxParser(
                file_type="docx", docx_include_headers=True
            ).parse_into_text(self.make_document()).content
        self.assertIn("Header identifier 3849", text)
        self.assertIn("Body text", text)

    def test_header_only_document_is_not_empty(self):
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Header only"
        output = io.BytesIO()
        document.save(output)
        text = Docx2Parser(
            file_type="docx", docx_include_headers=True
        ).parse_into_text(output.getvalue()).content
        self.assertIn("Header only", text)
