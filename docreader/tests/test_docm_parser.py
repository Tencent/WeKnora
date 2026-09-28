import io
import unittest
import zipfile
from unittest.mock import patch

from docx import Document as WordDocument
from lxml import etree
from PIL import Image

from docreader.models.document import Document
from docreader.parser.docm import normalize_docm_content_type
from docreader.parser.parser import Parser
from docreader.parser.registry import registry

DOCX_MAIN = (
    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)
DOCM_MAIN = b"application/vnd.ms-word.document.macroEnabled.main+xml"


def make_word_package(macro=True, content_types_encoding="UTF-8"):
    doc = WordDocument()
    doc.add_paragraph("DOCM content sentinel")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Group"
    table.cell(0, 1).text = "Item"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Merged category"
    table.cell(1, 1).text = "Alpha"
    table.cell(2, 1).text = "Beta"
    image = io.BytesIO()
    Image.new("RGB", (128, 96), "red").save(image, format="PNG")
    image.seek(0)
    doc.add_picture(image)
    output = io.BytesIO()
    doc.save(output)
    if not macro:
        return output.getvalue()
    converted = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as src, zipfile.ZipFile(
        converted, "w"
    ) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "[Content_Types].xml":
                data = data.replace(DOCX_MAIN, DOCM_MAIN)
                data = etree.tostring(
                    etree.fromstring(data),
                    encoding=content_types_encoding,
                    xml_declaration=True,
                )
            dst.writestr(info, data)
        # Inert opaque part: verifies preservation, not real VBA execution.
        dst.writestr("word/vbaProject.bin", b"inert macro preservation marker")
    return converted.getvalue()


class DocmParserTest(unittest.TestCase):
    def test_both_engine_entries_extract_text_tables_and_images(self):
        data = make_word_package()
        for engine in ("builtin", "markitdown"):
            with self.subTest(engine=engine):
                result = Parser().parse_file(
                    "sample.docm", ".DOCM", data, parser_engine=engine
                )
                self.assertIn("DOCM content sentinel", result.content)
                for item in ("Alpha", "Beta"):
                    row = next(
                        line for line in result.content.splitlines() if item in line
                    )
                    self.assertIn("Merged category", row)
                self.assertTrue(result.images)
                for ref in result.images:
                    self.assertIn(ref, result.content)
                listed = next(e for e in registry.list_engines() if e["name"] == engine)
                self.assertIn("docm", listed["file_types"])

    def test_builtin_fallback_reads_macro_content_type(self):
        with patch(
            "docreader.parser.markitdown_parser.MarkitdownParser.parse_into_text",
            return_value=Document(),
        ):
            result = Parser().parse_file(
                "sample.docm", "docm", make_word_package(), parser_engine="builtin"
            )
        self.assertIn("DOCM content sentinel", result.content)
        self.assertIn("Alpha", result.content)
        self.assertIn("Beta", result.content)
        self.assertTrue(result.images)

    def test_builtin_fallback_reads_utf16_content_types(self):
        data = make_word_package(content_types_encoding="UTF-16")
        with patch(
            "docreader.parser.markitdown_parser.MarkitdownParser.parse_into_text",
            return_value=Document(),
        ):
            result = Parser().parse_file(
                "sample.docm", "docm", data, parser_engine="builtin"
            )
        self.assertIn("DOCM content sentinel", result.content)
        self.assertIn("Alpha", result.content)
        self.assertIn("Beta", result.content)
        self.assertTrue(result.images)

    def test_normalization_preserves_every_other_package_member(self):
        data = make_word_package()
        normalized = normalize_docm_content_type(data)
        with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(
            io.BytesIO(normalized)
        ) as converted:
            self.assertEqual(original.namelist(), converted.namelist())
            for name in original.namelist():
                if name != "[Content_Types].xml":
                    self.assertEqual(original.read(name), converted.read(name))
            self.assertIn(DOCM_MAIN, original.read("[Content_Types].xml"))
            self.assertIn(DOCX_MAIN, converted.read("[Content_Types].xml"))
        self.assertEqual(
            WordDocument(io.BytesIO(normalized)).paragraphs[0].text,
            "DOCM content sentinel",
        )

    def test_docx_and_non_zip_are_unchanged(self):
        for data in (make_word_package(macro=False), b"not a word package"):
            self.assertIs(normalize_docm_content_type(data), data)


if __name__ == "__main__":
    unittest.main()
