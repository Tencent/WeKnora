import base64
import io
import logging
import shutil
import sys
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

# Avoid importing unrelated OCR/PDF engines from parser/__init__.py.
package = types.ModuleType("docreader.parser")
package.__path__ = [str(Path(__file__).resolve().parents[1] / "parser")]
sys.modules.setdefault("docreader.parser", package)

from docreader.parser.markdown_parser import MarkdownImageBase64
from docreader.parser.markitdown_parser import MarkitdownParser
from docreader.parser import embedded_image

FIXTURES = Path(__file__).with_name("fixtures")


def png(size=(16, 16)):
    output = io.BytesIO()
    Image.new("RGB", size, "red").save(output, format="PNG")
    return output.getvalue()


def inline(mime, data, title="diagram"):
    return f"![{title}](data:{mime};base64,{base64.b64encode(data).decode()})"


def word_with_preview(extension, data, mime, *, ole=False):
    """A real DOCX with a raster photo and a legacy VML/Office preview.

    The small EMF/WMF fixtures were exported from a 160x80 SVG containing a
    red rectangle and 'Diagram' text using LibreOffice Draw's emf/wmf filters.
    """
    document = WordDocument()
    document.add_paragraph("before preview")
    document.add_picture(io.BytesIO(png()))
    package = document.part.package
    part = Part(PackURI(f"/word/media/preview.{extension}"), mime, data, package)
    relation = document.part.relate_to(part, RELATIONSHIP_TYPE.IMAGE)
    container = OxmlElement("w:object" if ole else "w:pict")
    shape = etree.SubElement(container, "{urn:schemas-microsoft-com:vml}shape")
    image = etree.SubElement(shape, "{urn:schemas-microsoft-com:vml}imagedata")
    image.set(qn("r:id"), relation)
    if ole:
        original = Part(
            PackURI("/word/embeddings/object.bin"), "application/octet-stream",
            b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1embedded object", package,
        )
        object_relation = document.part.relate_to(original, RELATIONSHIP_TYPE.OLE_OBJECT)
        node = etree.SubElement(container, "{urn:schemas-microsoft-com:office:office}OLEObject")
        node.set(qn("r:id"), object_relation)
    document.add_paragraph().add_run()._r.append(container)
    document.add_paragraph("after preview")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class EmbeddedImageTest(unittest.TestCase):
    def parse(self, markdown):
        return MarkdownImageBase64().parse_into_text(markdown.encode())

    def test_ole_mislabeled_as_png_is_not_an_image(self):
        document = self.parse(inline("image/png", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1object"))
        self.assertEqual(document.images, {})
        self.assertEqual(document.content, "diagram")

    def test_visio_and_octet_stream_are_not_image_references(self):
        for mime in ("application/octet-stream", "application/vnd.ms-visio.drawing", "image/png"):
            with self.subTest(mime=mime):
                document = self.parse(inline(mime, b"PK\x03\x04visio"))
                self.assertEqual(document.images, {})
                self.assertEqual(document.content, "diagram")

    def test_failed_vector_does_not_remove_surrounding_text_or_good_picture(self):
        with (
            patch.object(embedded_image, "rasterize_media_bytes", return_value=None),
            patch.object(embedded_image, "_rasterize_with_libreoffice", return_value=None),
        ):
            document = self.parse(
                "before " + inline("image/x-emf", b"invalid metafile")
                + " after " + inline("image/png", png(), "photo")
            )
        self.assertEqual(len(document.images), 1)
        self.assertIn("before diagram after ![photo](images/", document.content)
        self.assertEqual(base64.b64decode(next(iter(document.images.values()))), png())

    def test_valid_raster_bytes_determine_extension_even_with_wrong_mime(self):
        for mime in ("image/png", "image/emf", "application/octet-stream"):
            with self.subTest(mime=mime):
                document = self.parse(inline(mime, png()))
                self.assertEqual(len(document.images), 1)
                self.assertTrue(next(iter(document.images)).endswith(".png"))
                self.assertEqual(base64.b64decode(next(iter(document.images.values()))), png())

    def test_docx_ole_without_renderable_preview_keeps_text_and_photo(self):
        raw = word_with_preview("bin", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1object", "image/png", ole=True)
        document = MarkitdownParser(file_type="docx").parse(raw)
        self.assertIn("before preview", document.content)
        self.assertIn("after preview", document.content)
        self.assertEqual(len(document.images), 1)
        self.assertEqual(base64.b64decode(next(iter(document.images.values()))), png())

    def test_missing_converters_skip_vector_preview(self):
        raw = word_with_preview("wmf", (FIXTURES / "embedded_diagram.wmf").read_bytes(), "image/wmf", ole=True)
        with (
            patch.object(embedded_image, "rasterize_media_bytes", return_value=None),
            patch.object(embedded_image, "find_soffice", return_value=None),
        ):
            document = MarkitdownParser(file_type="docx").parse(raw)
        self.assertEqual(len(document.images), 1)
        self.assertNotIn(".wmf", document.content)
        self.assertIn("after preview", document.content)

    def test_vector_bytes_are_converted_even_when_mime_is_wrong(self):
        raw = (FIXTURES / "embedded_diagram.emf").read_bytes()
        for mime in ("image/emf", "image/png", "application/octet-stream"):
            with self.subTest(mime=mime), patch.object(
                embedded_image, "rasterize_media_bytes", return_value=png()
            ) as convert:
                document = self.parse(inline(mime, raw))
                self.assertEqual(len(document.images), 1)
                self.assertTrue(next(iter(document.images)).endswith(".png"))
                self.assertEqual(base64.b64decode(next(iter(document.images.values()))), png())
                convert.assert_called_once_with("image.emf", raw)

    def test_native_docx_fallback_also_normalizes_vectors(self):
        from docreader.parser.docx_parser import Docx, _extract_image_in_process

        document = WordDocument()
        document.add_picture(io.BytesIO(png()))
        paragraph = document.paragraphs[0]
        relationship = paragraph._element.xpath(".//a:blip/@r:embed")[0]
        part = document.part.related_parts[relationship]
        part._blob = (FIXTURES / "embedded_diagram.wmf").read_bytes()
        part._content_type = "image/wmf"
        logger = logging.getLogger(__name__)
        with patch.object(embedded_image, "rasterize_media_bytes", return_value=png((160, 80))):
            image = Docx().get_picture(document, paragraph)
            self.assertEqual(image.size, (160, 80))
            image = _extract_image_in_process(logger, document, paragraph, 0, 0, 1920)
            self.assertEqual(image.size, (160, 80))

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice is not installed")
    def test_real_docx_vectors_and_ole_previews_become_decodable_png(self):
        # Force the default Docker path, which has LibreOffice, not ImageMagick.
        for extension in ("wmf", "emf"):
            with self.subTest(extension=extension), patch.object(
                embedded_image, "rasterize_media_bytes", return_value=None
            ):
                raw = word_with_preview(
                    extension, (FIXTURES / f"embedded_diagram.{extension}").read_bytes(),
                    f"image/{extension}", ole=True,
                )
                document = MarkitdownParser(file_type="docx").parse(raw)
                self.assertEqual(len(document.images), 2)
                rendered_preview = False
                for name, encoded in document.images.items():
                    self.assertTrue(name.endswith(".png"))
                    with Image.open(io.BytesIO(base64.b64decode(encoded))) as image:
                        image.load()
                        self.assertEqual(image.format, "PNG")
                        if image.width > 16:
                            colors = image.convert("RGB").getcolors(image.width * image.height)
                            red_pixels = sum(count for count, (r, g, b) in colors if r > 200 and g < 50 and b < 50)
                            self.assertGreater(red_pixels, image.width * image.height / 2)
                            rendered_preview = True
                self.assertTrue(rendered_preview)
                self.assertIn("before preview", document.content)
                self.assertIn("after preview", document.content)

    def test_metafile_conversion_timeout_skips_image(self):
        with (
            patch.object(embedded_image, "rasterize_media_bytes", return_value=None),
            patch.object(embedded_image, "find_soffice", return_value="soffice"),
            patch.object(embedded_image.subprocess, "run", side_effect=embedded_image.subprocess.TimeoutExpired("soffice", 60)),
        ):
            document = self.parse(inline("image/emf", b"metafile"))
        self.assertEqual(document.images, {})
        self.assertEqual(document.content, "diagram")


if __name__ == "__main__":
    unittest.main()
